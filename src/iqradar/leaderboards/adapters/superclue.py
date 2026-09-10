from __future__ import annotations

import io
import json
import logging
import urllib.request
import urllib.parse
from datetime import date, datetime, timezone
from pathlib import Path

from iqradar.leaderboards.base import LeaderboardSource
from iqradar.leaderboards.schemas import LeaderboardEntry, SourceConfig

DEFAULT_BASE_URL = "https://www.superclueai.com"
FETCH_TIMEOUT_SEC = 20

# xlsx 文件名模式: /data/generalboard/{YYYY年M月}.xlsx
# 例如: 2026年7月.xlsx

# sheet 名称映射（sheet 名 → benchmark 标签）
SHEET_MAP = {
    "总排行榜": "总排行榜",
    "推理模型总排行榜": "推理模型总排行榜",
    "推理任务总排行榜": "推理任务总排行榜",
    "开源排行榜": "开源排行榜",
}

# 错误容忍：缺列时跳过
_MIN_COLS = 4


class SuperCLUESource(LeaderboardSource):
    """SuperCLUE 中文通用大模型综合基准（每月更新，xlsx 下载）。"""

    def __init__(self, config: SourceConfig, cache_root: Path) -> None:
        super().__init__(config, cache_root)
        self._base_url = (config.base_url or DEFAULT_BASE_URL).rstrip("/")
        self._extra = config.extra or {}
        # 允许通过 extra.month 覆盖（如 "2026年7月"），默认取当前月
        self._month: str = self._extra.get("month") or _current_month()

    @property
    def display_name(self) -> str:
        return "SuperCLUE"

    async def fetch(self) -> list[LeaderboardEntry]:
        try:
            content = self._download_xlsx()
        except Exception as exc:
            logging.warning("SuperCLUE xlsx download failed: %s", exc)
            return []

        try:
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        except Exception as exc:
            logging.warning("SuperCLUE xlsx parse failed: %s", exc)
            return []

        entries: list[LeaderboardEntry] = []
        now = datetime.now(timezone.utc)

        for sheet_name, bench_label in SHEET_MAP.items():
            if sheet_name not in workbook.sheetnames:
                continue
            try:
                rows = self._parse_sheet(workbook[sheet_name])
                entries.extend(
                    self._to_entries(sheet_name, bench_label, rows, now)
                )
            except Exception as exc:
                logging.warning(
                    "SuperCLUE sheet %s parse failed: %s", sheet_name, exc
                )

        workbook.close()
        return entries

    _SHEET_COL_BENCHMARK: dict[str, tuple[str, ...]] = {
        # 总排行榜 / 推理模型总排行榜：rank, model, provider, open?, 总分, 数学推理, 科学推理,
        #                               代码生成, 智能体Agent, 精确指令遵循, 幻觉控制
        "总排行榜": ("总分", "数学推理", "科学推理", "代码生成", "智能体Agent", "精确指令遵循", "幻觉控制"),
        "推理模型总排行榜": ("总分", "数学推理", "科学推理", "代码生成", "智能体Agent", "精确指令遵循", "幻觉控制"),
        "推理任务总排行榜": ("总分", "数学推理", "科学推理", "代码生成", "智能体Agent"),
        "开源排行榜": ("总分", "数学推理", "科学推理", "代码生成", "智能体Agent", "精确指令遵循", "幻觉控制"),
    }

    def _parse_sheet(self, worksheet) -> list[tuple]:
        """读取 sheet，返回 (模型名, 供应商, 分数字典) 列表。"""
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c) if c is not None else "" for c in rows[0]]
        # 找到模型名列（通常是第 2 列）与供应商列（第 3 列）
        model_col = _find_col(header, ("模型", "Model", "model"))
        provider_col = _find_col(header, ("厂商", "公司", "机构", "组织", "provider"))

        result: list[tuple] = []
        for row in rows[1:]:
            if not row or len(row) < _MIN_COLS:
                continue
            model_name = str(row[model_col]).strip() if model_col is not None else ""
            if not model_name or model_name.startswith("排名"):
                continue
            provider = str(row[provider_col]).strip() if provider_col is not None else None
            scores: dict[str, float] = {}
            for i, cell in enumerate(row):
                if i >= len(header) or header[i] in ("", "排名", "序号"):
                    continue
                label = header[i]
                try:
                    value = _to_float(cell)
                except (TypeError, ValueError):
                    continue
                if value is not None:
                    scores[label] = value
            result.append((model_name, provider, scores))
        return result

    def _to_entries(
        self,
        sheet_name: str,
        bench_label: str,
        rows: list[tuple],
        now: datetime,
    ) -> list[LeaderboardEntry]:
        entries: list[LeaderboardEntry] = []
        sub_dims = self._SHEET_COL_BENCHMARK.get(sheet_name, ("总分",))
        for model_name, provider, scores in rows:
            if not scores:
                continue
            total = scores.get("总分")
            if total is None:
                # 取第一个可用维度
                for dim in sub_dims:
                    if dim in scores:
                        total = scores[dim]
                        break
            if total is None:
                continue

            # 各维度得分放 metadata
            dim_scores = {k: v for k, v in scores.items() if k in sub_dims}

            entries.append(
                LeaderboardEntry(
                    source_id=self.source_id,
                    model_name=model_name,
                    model_display=model_name,
                    provider=provider,
                    benchmark_name=self._bench_slug(sheet_name),
                    benchmark_label=bench_label,
                    score=float(total),
                    score_label="总分",
                    rank=None,
                    ci=None,
                    metadata={
                        "month": self._month,
                        "sheet": sheet_name,
                        "dimensions": dim_scores,
                    },
                    fetched_at=now,
                )
            )
        return entries

    @staticmethod
    def _bench_slug(sheet_name: str) -> str:
        if sheet_name == "总排行榜":
            return "general"
        if sheet_name == "推理模型总排行榜":
            return "reasoning_models"
        if sheet_name == "推理任务总排行榜":
            return "reasoning_tasks"
        if sheet_name == "开源排行榜":
            return "open_source"
        return "general"

    def _download_xlsx(self) -> bytes:
        """下载本月 xlsx 文件。"""
        filename = f"{self._month}.xlsx"
        url = f"{self._base_url}/data/generalboard/{urllib.parse.quote(filename)}"
        req = urllib.request.Request(
            url, method="GET", headers={"User-Agent": "IQRadar/1.0"}
        )
        with urllib.request.urlopen(req, timeout=FETCH_TIMEOUT_SEC) as resp:
            return resp.read()


def _current_month() -> str:
    today = date.today()
    # 中文月份，如 "2026年7月"
    return f"{today.year}年{today.month}月"


def _find_col(header: list[str], names: tuple[str, ...]) -> int | None:
    for i, label in enumerate(header):
        for name in names:
            if name.lower() in label.lower():
                return i
    return None


def _to_float(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None